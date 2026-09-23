"""Import explicitly supplied local billing CSV/review XLSX files, without network access.

python -m back.api.import_data --details /path/details_long.csv --review /path/review.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from . import catalog, db


DETAIL_FIELDS = {
    "총사용금액": "total_use_amount_won", "차감금액": "deduct_amount_won", "납부금액": "pay_amount_won",
    "상수도_기본료": "water_base_won", "상수도_사용료": "water_use_won", "하수도_사용료": "sewer_total_won",
    "물이용부담금": "waterload_total_won", "사용량": "usage", "당월지침": "reading_curr",
    "전월지침": "reading_prev", "전납기사용량": "prev_period_usage", "전년동기사용량": "prev_year_usage",
    "공동사용량": "common_usage", "총사용량": "total_usage", "체납액": "arrears_won", "미납액": "unpaid_won",
}


def upsert_bill_summaries(conn, rows):
    """Apply collected list rows without inventing detailed fees or deleting history."""
    counts = {"imported": 0, "skipped": 0}
    for row in rows:
        cid = catalog.customer_id(row.get("mkey") or row.get("고객번호"))
        raw_date = str(row.get("napgi_compact") or row.get("napgi") or row.get("납기일") or "")
        match = re.match(r"^(\d{4})\D?(0[1-9]|1[0-2])(?:\D?(\d{2}))?", raw_date)
        meter = conn.execute("SELECT id FROM meters WHERE provider='arisu' AND customer_number=?", (cid,)).fetchone()
        if not match or not meter:
            counts["skipped"] += 1
            continue
        period = f"{match[1]}-{match[2]}"
        if match[3]:
            try:
                datetime.strptime(f"{period}-{match[3]}", "%Y-%m-%d")
            except ValueError:
                counts["skipped"] += 1
                continue
        gubun = row.get("gubun") or row.get("구분") or "정기분"
        notice = str(row.get('notice_number') or row.get('고지번호') or '').strip()
        if notice and not re.fullmatch(r'\d{6,12}', notice):
            counts['skipped'] += 1
            continue
        old = conn.execute("SELECT payload,source FROM bills WHERE meter_id=? AND period=? AND gubun=? AND notice_number=?",
                           (meter["id"], period, gubun, notice)).fetchone()
        bill = json.loads(old["payload"]) if old else {}
        bill.update(ym=period, year=int(match[1]), month=int(match[2]), gubun=gubun)
        if notice:
            bill.update(notice_number=notice, 고지번호=notice)
        for target, source in (("부과금액", "bugwa_amount_won"), ("총사용량", "total_usage_ton")):
            value = catalog.number(row.get(source))
            if value is not None:
                bill[target] = value
        if row.get("sunap_status"):
            bill["수납상태"] = row["sunap_status"]
        if match[3]:
            bill["납기일"] = f"{period}-{match[3]}"
        details = row.get('details')
        if details is None or isinstance(details, str) and not details.strip():
            details = {}
        elif isinstance(details, str):
            try:
                details = json.loads(details)
            except json.JSONDecodeError:
                raise ValueError('청구 상세 자료 details의 JSON 형식이 올바르지 않습니다.') from None
        if not isinstance(details, dict):
            raise ValueError('청구 상세 자료 details는 JSON 객체여야 합니다.')
        detailed_values = {key: value for key in (*DETAIL_FIELDS, '상수도_합계', '지하수사용량', '지하수당월지침', '지하수전월지침', '가구수', '정기검침일', '계량기대금', '설치비', '연체금')
                           if (value := catalog.number(details.get(key))) is not None}
        detailed_values.update({key: str(details[key]) for key in ('periodStart', 'periodEnd', '납기일', '주소', '고지서성명',
                               '전자수용가번호', '전자납부번호', '업종', '계량기번호', '구경', '납부방법') if details.get(key)})
        bill.update(detailed_values)
        source = 'i121_public_detail' if detailed_values and row.get('detail_source') == 'i121_public_detail' else old['source'] if old else 'i121_summary'
        bill["detail_available"] = catalog.number(bill.get("납부금액")) is not None
        bill["summary_only"] = not (detailed_values or source in ('details_csv', 'i121_public_detail')
                                    or catalog.bill_usage_field(bill) == '사용량')
        bill["summary_updated_at"] = db.now()
        conn.execute("""INSERT INTO bills(id,meter_id,period,gubun,payload,source,notice_number) VALUES(?,?,?,?,?,?,?)
                       ON CONFLICT(meter_id,period,gubun,notice_number) DO UPDATE SET payload=excluded.payload,source=excluded.source""",
                     (catalog.bill_id(meter["id"], period, gubun, notice), meter["id"], period, gubun,
                      json.dumps(bill, ensure_ascii=False), source, notice))
        counts["imported"] += 1
    return counts


def _ensure_meter(conn, cid, row):
    existing = conn.execute("SELECT id FROM meters WHERE provider='arisu' AND customer_number=?", (cid,)).fetchone()
    if existing:
        return existing["id"]
    display = str(row.get("역명") or row.get("csNm") or cid).strip()
    name = re.sub(r"\d+$", "", display)
    office = str(row.get("사업소명") or row.get("영업사업소") or "미지정").strip()
    conn.execute("INSERT OR IGNORE INTO offices(id,name) VALUES(?,?)", (office, office))
    conn.execute("INSERT OR IGNORE INTO stations(id,name) VALUES(?,?)", (name, name))
    now = datetime.now(timezone.utc).isoformat()
    conn.execute("""INSERT INTO meters(id,provider,customer_number,station_id,office_id,line,display_name,
                 purpose,tariff,address,active,daily_enabled,metadata,created_at,updated_at)
                 VALUES(?,'arisu',?,?,?,?,?,?,?,?,1,0,'{}',?,?)""",
                 (cid, cid, name, office, str(row.get("호선") or ""), display,
                  str(row.get("용도표기") or ""), str(row.get("usage_type") or ""), str(row.get("addr") or ""), now, now))
    return cid


def import_details(conn, path: Path):
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        columns = set(csv.DictReader(stream).fieldnames or [])
    if not {"mkey", "napgi", "gubun", "pay_amount_won"}.issubset(columns):
        raise ValueError("상세 CSV에는 mkey, napgi, gubun, pay_amount_won 열이 필요합니다.")
    counts = {"imported": 0, "skipped": 0, "missing_details": 0}
    for row in catalog.read_csv(path):
        cid, period = catalog.customer_id(row.get("mkey")), row.get("napgi", "")
        if not re.fullmatch(r"\d{9}", cid) or not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", period):
            counts["skipped"] += 1
            continue
        if catalog.number(row.get("pay_amount_won")) is None:
            counts["missing_details"] += 1
        meter_id = _ensure_meter(conn, cid, row)
        gubun = row.get("gubun") or "정기분"
        notice = str(row.get('notice_number') or row.get('고지번호') or '').strip()
        if notice and not re.fullmatch(r'\d{6,12}', notice):
            counts['skipped'] += 1
            continue
        old = conn.execute("SELECT payload FROM bills WHERE meter_id=? AND period=? AND gubun=? AND notice_number=?",
                           (meter_id, period, gubun, notice)).fetchone()
        bill = json.loads(old["payload"]) if old else {}
        bill.update(ym=period, year=int(period[:4]), month=int(period[5:]), gubun=gubun)
        if notice:
            bill.update(notice_number=notice, 고지번호=notice)
        for target, source in DETAIL_FIELDS.items():
            value = catalog.number(row.get(source))
            if value is not None or target not in bill:
                bill[target] = value
        for target, source in {"periodStart": "use_period_start", "periodEnd": "use_period_end",
                               "납기일": "due_date", "수납상태": "sunap_status", "주소": "addr",
                               "고지서성명": "bill_name_masked", "전자수용가번호": "e_customer_no",
                               "전자납부번호": "e_payment_no", "계량기번호": "meter_no", "구경": "meter_size"}.items():
            if row.get(source):
                bill[target] = row[source]
        bill["detail_available"] = catalog.number(bill.get("납부금액")) is not None
        bill["summary_only"] = False
        conn.execute("""INSERT INTO bills(id,meter_id,period,gubun,payload,source,notice_number) VALUES(?,?,?,?,?,?,?)
                       ON CONFLICT(meter_id,period,gubun,notice_number) DO UPDATE SET payload=excluded.payload,source=excluded.source""",
                     (catalog.bill_id(meter_id, period, gubun, notice), meter_id, period, gubun,
                      json.dumps(bill, ensure_ascii=False), "details_csv", notice))
        meter = conn.execute("SELECT metadata FROM meters WHERE id=?", (meter_id,)).fetchone()
        meta = json.loads(meter["metadata"] or "{}")
        if not meta.get('arisu_customer_name') and row.get('csNm'):
            meta['arisu_customer_name'] = re.sub(r'\s+', ' ', row['csNm']).strip()
        for target, source in {"계량기번호": "meter_no", "구경": "meter_size", "전자수용가번호": "e_customer_no",
                               "전자납부번호": "e_payment_no", "가구수": "households"}.items():
            if not meta.get(target) and row.get(source):
                meta[target] = row[source]
        conn.execute("UPDATE meters SET metadata=?,purpose=CASE WHEN purpose='' THEN ? ELSE purpose END WHERE id=?",
                     (json.dumps(meta, ensure_ascii=False), row.get("용도표기") or "", meter_id))
        counts["imported"] += 1
    return counts


def import_review(conn, path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["수도요금 고지서 정보"] if "수도요금 고지서 정보" in workbook.sheetnames else workbook.active
    headers = None
    count = 0
    try:
        for values in sheet.iter_rows(values_only=True):
            normalized = [str(v or "").strip().replace(" ", "").replace("\n", "") for v in values]
            if headers is None:
                if "고객번호" in normalized:
                    headers = normalized
                continue
            row = {h: v for h, v in zip(headers, values) if h and v is not None}
            cid = catalog.customer_id(row.get("고객번호"))
            if not re.fullmatch(r"\d{9}", cid):
                continue
            row["용도표기"] = row.get("용도표기") or ""
            mid = _ensure_meter(conn, cid, row)
            stored = conn.execute("SELECT metadata,purpose,daily_enabled FROM meters WHERE id=?", (mid,)).fetchone()
            meta = json.loads(stored["metadata"] or "{}")
            notes = {key: str(value) for key, value in row.items() if "검토" in key}
            meta["review_notes"] = notes
            if row.get('고지서상성명'):
                meta['arisu_customer_name'] = re.sub(r'\s+', ' ', str(row['고지서상성명'])).strip()
            daily_enabled = stored['daily_enabled']
            if '유무' in headers:
                mark = str(row.get('유무') or '').strip().upper()
                if mark not in ('', 'O'):
                    raise ValueError('검토 파일의 유무 열은 O 또는 빈칸이어야 합니다.')
                daily_enabled = int(mark == 'O')
            conn.execute("UPDATE meters SET metadata=?,daily_enabled=?,purpose=CASE WHEN purpose='' THEN ? ELSE purpose END WHERE id=?",
                         (json.dumps(meta, ensure_ascii=False), daily_enabled, str(row.get("용도표기") or ""), mid))
            count += 1
        if headers is None:
            raise ValueError("검토 파일에서 고객 번호 열을 찾을 수 없습니다.")
    finally:
        workbook.close()
    return {"imported": count}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--details", type=Path)
    parser.add_argument("--review", type=Path)
    args = parser.parse_args()
    if not args.details and not args.review:
        parser.error("--details 또는 --review 파일을 지정하세요.")
    for path in (args.details, args.review):
        if path and not path.is_file():
            parser.error(f"파일이 없습니다: {path}")
    db.init_db()
    results = {}
    with db.connect() as conn:
        catalog.bootstrap(conn)
        if args.details:
            results["details"] = import_details(conn, args.details)
        if args.review:
            results["review"] = import_review(conn, args.review)
    print(json.dumps(results, ensure_ascii=False))


if __name__ == "__main__":
    main()
