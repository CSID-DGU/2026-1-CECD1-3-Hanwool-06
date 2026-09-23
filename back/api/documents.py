"""Generate private billing documents from stored values, never guessed amounts."""
from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from . import catalog


FONT = "NanumGothic"
if FONT not in pdfmetrics.getRegisteredFontNames():
    pdfmetrics.registerFont(TTFont(FONT, str(Path(__file__).parent / "fonts/NanumGothic-Regular.ttf")))


def _text(value):
    return "-" if value is None or value == "" else str(value)


def _amount(value):
    if value is None or value == "":
        return "-"
    try:
        return f"{float(value):,.0f}"
    except (ValueError, TypeError):
        return _text(value)


def bill_pdf(meter: dict, bill: dict) -> bytes:
    """A4 statement; source-labelled and safe for long Korean addresses/identifiers."""
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm,
                           topMargin=18 * mm, bottomMargin=18 * mm, title="수도요금 조회내역")
    normal = ParagraphStyle("normal", fontName=FONT, fontSize=9, leading=14, wordWrap="CJK")
    title = ParagraphStyle("title", parent=normal, fontSize=20, leading=28, textColor=colors.HexColor("#17345c"))
    small = ParagraphStyle("small", parent=normal, fontSize=8, leading=12, textColor=colors.HexColor("#5d6776"))

    def p(value, style=normal):
        return Paragraph(escape(_text(value)).replace("\n", "<br/>"), style)

    def table(rows, widths, header=False):
        item = Table([[p(c) for c in row] for row in rows], colWidths=[w * mm for w in widths], repeatRows=int(header))
        commands = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 9),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LINEBELOW", (0, 0), (-1, -1), .4, colors.HexColor("#d8e0e9"))]
        if header:
            commands.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf2f8")))
        item.setStyle(TableStyle(commands))
        return item

    metadata = meter.get("metadata") or {}
    if isinstance(metadata, str):
        metadata = json.loads(metadata)
    flow = [p("수도요금 조회내역", title), p(f"{bill.get('ym', '')}  {bill.get('gubun', '정기분')}", small), Spacer(1, 7 * mm)]
    if bill.get('notice_number') or bill.get('고지번호'):
        flow.append(p('고지번호: ' + str(bill.get('notice_number') or bill['고지번호'])))
    flow.append(table([
        ["역 / 계량기", meter.get("display_name") or meter.get("station_name"), "고객번호", meter.get("customer_number")],
        ["영업사업소", meter.get("office_name"), "사용 목적", meter.get("purpose")],
        ["주소", bill.get("주소") or meter.get("address"), "요금 업종", bill.get("업종") or bill.get("용도") or meter.get("tariff")],
        ["고지서상 성명", bill.get("고지서성명") or bill.get("사용자명") or bill.get("성명"), "수납상태", bill.get("수납상태")],
        ["사용기간", f"{_text(bill.get('periodStart'))} ~ {_text(bill.get('periodEnd'))}", "납기일", bill.get("납기일")],
        ["계량기번호", bill.get("계량기번호") or metadata.get("계량기번호"), "구경", bill.get("구경") or metadata.get("구경")],
        ["전자수용가번호", bill.get("전자수용가번호") or metadata.get("전자수용가번호"),
         "전자납부번호", bill.get("전자납부번호") or metadata.get("전자납부번호")],
    ], [32, 55, 32, 55]))
    amount_label = "납부금액" if bill.get("납부금액") is not None else "부과금액"
    flow += [Spacer(1, 7 * mm), p(amount_label + "  " + _amount(bill.get(amount_label)) + " 원", title), Spacer(1, 5 * mm)]
    flow.append(table([
        ["총 사용금액 (원)", "차감금액 (원)", amount_label + " (원)"],
        [_amount(bill.get(k)) for k in ("총사용금액", "차감금액", amount_label)],
    ], [58, 58, 58], True))
    additional_fees = [key for key in ("계량기대금", "설치비", "연체금") if bill.get(key) is not None]
    if additional_fees:
        flow += [Spacer(1, 3 * mm), table([
            [key + " (원)" for key in additional_fees],
            [_amount(bill[key]) for key in additional_fees],
        ], [174 / len(additional_fees)] * len(additional_fees), True)]
    usage = catalog.bill_usage(bill)
    usage_label = "상·하수도 사용량 (톤)" if catalog.bill_usage_field(bill) == "사용량" else "총사용량 (톤)"
    flow += [Spacer(1, 5 * mm), table([
        ["내역", "금액 (원)", "검침 / 사용량", "값"],
        ["상수도 기본요금", _amount(bill.get("상수도_기본료")), "전월 지침", _amount(bill.get("전월지침"))],
        ["상수도 사용요금", _amount(bill.get("상수도_사용료")), "당월 지침", _amount(bill.get("당월지침"))],
        ["하수도요금", _amount(bill.get("하수도_사용료")), usage_label, _amount(usage)],
        ["물이용부담금", _amount(bill.get("물이용부담금")), "전납기 사용량 (톤)", _amount(bill.get("전납기사용량"))],
        ["수납상태", bill.get("수납상태"), "전년 동기 사용량 (톤)", _amount(bill.get("전년동기사용량"))],
    ], [48, 39, 48, 39], True)]
    groundwater = [bill.get(key) if bill.get(key) is not None else bill.get(old) for key, old in (
        ("지하수전월지침", "지하수_전월지침"), ("지하수당월지침", "지하수_당월지침"), ("지하수사용량", "지하수_사용량"))]
    if any(value is not None for value in groundwater):
        flow += [Spacer(1, 5 * mm), table([
            ["지하수 전월지침", "지하수 당월지침", "지하수 사용량 (톤)"],
            [_amount(value) for value in groundwater],
        ], [58, 58, 58], True)]
    source = {"legacy_seed": "기존 청구자료", "details_csv": "아리수 상세 수집자료",
              "i121_summary": "아리수 청구목록 (상세 내역 미수집)",
              "i121_public_detail": "아리수 요금 전체조회 상세자료"}.get(bill.get("source"), "등록 청구자료")
    flow += [Spacer(1, 3 * mm), p("수집된 청구 데이터를 바탕으로 생성한 조회내역입니다. 아리수 원본 고지서가 아닙니다.", small),
             p("자료가 없는 항목은 -로 표시합니다.  /  자료 출처: " + source, small)]

    def footer(canvas, document):
        canvas.setFont(FONT, 8)
        canvas.setFillColor(colors.HexColor("#5d6776"))
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, str(document.page))

    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export_xlsx(rows: list[dict]) -> bytes:
    """Export server-authorized rows. External strings remain text, never formulas."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "조회내역"
    columns = list(dict.fromkeys(key for row in rows for key in row))
    sheet.append(columns or ["조회 결과 없음"])
    for row in rows:
        sheet.append([json.dumps(row[key], ensure_ascii=False) if isinstance(row.get(key), (dict, list))
                      else row.get(key) for key in columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    # Column labels can also originate from an uploaded source.
    for cell in sheet[1]:
        cell.data_type = "s"
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17345C")
    for column in sheet.columns:
        width = max((len(str(c.value or "")) for c in column), default=12)
        sheet.column_dimensions[column[0].column_letter].width = min(50, max(16, width + 2))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
