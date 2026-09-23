import csv
import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from back.api import catalog, config, data_access, db, documents, import_data


class CatalogCheck(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.seed = self.root / "seed"
        self.seed.mkdir()
        bill = {"ym": "2026-04", "year": 2026, "month": 4, "납부금액": 1234, "사용량": 10,
                "periodStart": "2026-02-01", "periodEnd": "2026-03-31"}
        self.seed.joinpath("bills.json").write_text(json.dumps({
            "000000001": {"역명": "가역1", "사업소명": "가사업소", "주소": "서울", "용도": "일반용", "bills": [bill]},
            "000000002": {"역명": "가역2", "사업소명": "나사업소", "주소": "서울", "용도": "일반용", "bills": []},
        }, ensure_ascii=False))
        self.seed.joinpath("stations.json").write_text(json.dumps({
            "000000001": {"역명": "가역1", "호선": 1, "영업사업소": "가사업소"},
            "000000002": {"역명": "가역2", "호선": 2, "영업사업소": "나사업소"},
        }, ensure_ascii=False))
        self.seed.joinpath("locations.json").write_text('{"000000001":{"x":20,"y":30}}')
        self.patch = patch.multiple(config, ROOT=self.root, SEED_DIR=self.seed,
                                    DATA_DIR=self.root / "runtime", APP_DB_PATH=self.root / "app.sqlite3")
        self.patch.start()
        db.init_db()
        with db.connect() as conn:
            catalog.bootstrap(conn)
            conn.execute('UPDATE meters SET daily_enabled=1')

    def tearDown(self):
        self.patch.stop()
        self.tmp.cleanup()

    def write_csv(self, path, rows):
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)

    def test_scope_and_snapshot(self):
        version = config.DATA_DIR / "snapshot/versions/check"
        version.mkdir(parents=True)
        series = {"000000001": {"usage": [{"date": "2026-05-01", "value": 10}], "ridership": []},
                  "000000002": {"usage": [{"date": "2026-06-01", "value": 90}], "ridership": []}}
        (version / "daily.json").write_text(json.dumps(series))
        (version / "risk.json").write_text("{}")
        (version.parent.parent / "manifest.json").write_text('{"published":true,"directory":"versions/check"}')
        with db.connect() as conn:
            conn.execute("UPDATE meters SET purpose='직원용' WHERE id='000000001'")
            catalog.bootstrap(conn)
            self.assertEqual(len(catalog.list_meters(conn)), 2)
        data = catalog.payload(["가사업소"])
        self.assertEqual(set(data["bills"]), {"000000001"})
        self.assertEqual(data["status"]["latest_usage"], "2026-05-01")
        self.assertEqual(data["meters"][0]["purpose"], "직원용")
        self.assertEqual(data["meters"][0]["data_mode"], "daily")
        self.assertEqual(data["status"]["counts"], {"daily": 1, "billing_only": 0, "pending": 0})
        self.assertEqual(set(data["locations"]), {"가역"})
        self.assertEqual(catalog.payload([])["meters"], [])
        self.assertIsNone(catalog.payload([])["status"]["reference_date"])
        self.assertEqual(catalog.payload([])["status"]["counts"], {"daily": 0, "billing_only": 0, "pending": 0})
        self.assertEqual(catalog.payload(["나사업소"])["bills"]["000000002"]["bills"], [])

    def test_merge_details_and_risk(self):
        path = self.root / "details.csv"
        self.write_csv(path, [
            {"mkey": "1", "napgi": "2026-05", "gubun": kind, "pay_amount_won": "2300",
             "sewer_total_won": "700", "waterload_total_won": "90", "usage": ""}
            for kind in ("정기분", "수시분")])
        with db.connect() as conn:
            self.assertEqual(import_data.import_details(conn, path)["imported"], 2)
            import_data.import_details(conn, path)
            self.assertEqual(conn.execute("SELECT count(*) FROM bills").fetchone()[0], 3)
            self.write_csv(path, [{"mkey": "1", "napgi": "2026-06", "gubun": "수시분", "pay_amount_won": ""}])
            self.assertEqual(import_data.import_details(conn, path)["missing_details"], 1)
        seed_risk = self.seed / 'risk/test_anomalies.csv'
        archived_risk = self.root / "analysis/risk_snapshot/scripts/LightGBM_Model/results/test_anomalies.csv"
        self.write_csv(seed_risk, [{"고객번호": "1", "날짜": "2026-06-13", "일사용량_톤": "10", "pred_ton": "7",
                            "residual": "3", "z_score": "4", "심각도": "주의", "방향": "과다", "data_error_candidate": "False"}])
        self.write_csv(archived_risk, [{"고객번호": "1", "날짜": "2026-07-01", "일사용량_톤": "10", "predicted_ton": "8",
                            "error_ton": "2", "deviation_score": "3", "심각도": "정상", "방향": "", "likely_data_error": "False"}])
        self.write_csv(self.root / "data/processed/daily_water_usage.csv", [
            {"고객번호": "1", "역명": "가역1", "검침일": "2026-06-13", "일사용량_톤": "10"}])
        data = catalog.payload(["가사업소"])
        self.assertEqual(data["status"]["latest_risk"], "2026-06-13")
        self.assertEqual(data["risk"]["000000001"][-1]["pred"], 7)
        bills = data["bills"]["000000001"]["bills"]
        self.assertEqual({b["gubun"] for b in bills if b["ym"] == "2026-05"}, {"정기분", "수시분"})
        self.assertIsNone(bills[-1]["사용량"])
        self.assertFalse(bills[-1]["detail_available"])

    def test_archived_meter_stays_private_and_restores_preserved_history(self):
        deleted_at = "2026-09-23T00:00:00+00:00"
        with db.connect() as conn:
            original_bill = dict(conn.execute("SELECT * FROM bills WHERE meter_id='000000001'").fetchone())
            conn.execute("UPDATE meters SET deleted_at=?,active=0,daily_enabled=1 WHERE id='000000001'", (deleted_at,))
            # Even a deliberate seed replay must not resurrect an archived contract.
            conn.execute("DELETE FROM settings WHERE key='catalog_seed_v1'")
            catalog.bootstrap(conn)
            self.assertEqual([m['id'] for m in catalog.list_meters(conn)], ['000000002'])
            self.assertEqual(catalog.list_meters(conn, ['가사업소']), [])
            archived = catalog.list_meters(conn, ['가사업소'], include_deleted=True)
            self.assertEqual([m['id'] for m in archived], ['000000001'])
            self.assertEqual(archived[0]['deleted_at'], deleted_at)
            self.assertEqual(archived[0]['active'], 0)
            self.assertEqual([m['id'] for m in catalog.list_meters(conn, ['나사업소'], include_deleted=True)], ['000000002'])
            self.assertEqual(catalog.list_meters(conn, [], include_deleted=True), [])
            self.assertEqual(dict(conn.execute("SELECT * FROM bills WHERE meter_id='000000001'").fetchone()), original_bill)
        self.assertEqual([m['id'] for m in catalog.collector_meters()], ['000000002'])
        daily = {cid: {'usage': [{'date': date, 'value': 10}], 'ridership': []}
                 for cid, date in [('000000001', '2026-09-22'), ('000000002', '2026-08-31')]}
        risk = {cid: [{'date': rows['usage'][0]['date'], 'actual': 10, 'pred': 5, 'residual': 5,
                       'z': 4, 'severity': '경고'}] for cid, rows in daily.items()}
        with patch.object(catalog, 'data_sources', return_value=(daily, risk)), patch.object(config, 'TODAY_OVERRIDE', ''):
            payload = catalog.payload(['가사업소'])
            for field in ('bills', 'stations', 'daily', 'risk', 'locations'):
                self.assertEqual(payload[field], {})
            self.assertEqual(payload['meters'], [])
            self.assertIsNone(payload['status']['reference_date'])
            self.assertEqual(payload['status']['counts'], {'daily': 0, 'billing_only': 0, 'pending': 0})
            self.assertEqual(data_access.reference_date(), '2026-08-31')
            self.assertEqual(data_access.anomalies_on('2026-09-22'), [])
            self.assertEqual(data_access.station_history('000000001', '2026-09-22'), [])
            with db.connect() as conn:
                conn.execute("UPDATE meters SET deleted_at=NULL,active=1 WHERE id='000000001'")
            restored = catalog.payload(['가사업소'])
            self.assertEqual(restored['bills']['000000001']['bills'][0]['id'], original_bill['id'])
            self.assertEqual(restored['bills']['000000001']['bills'][0]['납부금액'], 1234)
            self.assertEqual(set(restored['risk']), {'000000001'})
            self.assertEqual(data_access.reference_date(), '2026-09-22')
            self.assertEqual([m['id'] for m in catalog.collector_meters()], ['000000001', '000000002'])

    def test_documents(self):
        from openpyxl import load_workbook
        from pypdf import PdfReader

        meter = {"display_name": "서울역 지하 공동사용 계량기", "customer_number": "000000001", "office_name": "가사업소",
                 "address": "서울특별시 종로구 긴 주소 지하 연결통로 시설관리실 " * 3, "purpose": "시민용", "tariff": "일반용"}
        bill = {"ym": "2026-05", "납부금액": 1234, "source": "details_csv", "notice_number": "900000001"}
        pdf = documents.bill_pdf(meter, bill)
        reader = PdfReader(BytesIO(pdf))
        self.assertEqual(len(reader.pages), 1)
        self.assertIn("1,234", reader.pages[0].extract_text())
        self.assertIn("아리수 원본", reader.pages[0].extract_text())
        self.assertIn('900000001', reader.pages[0].extract_text())
        self.assertTrue(any(font.get_object().get("/FontDescriptor", {}).get("/FontFile2")
                            for font in reader.pages[0]["/Resources"]["/Font"].values()))
        content = documents.export_xlsx([{"고객번호": "000000001", "이름": "=HYPERLINK(\"bad\")", "값": 0},
                                         {"고객번호": "000000002", "이름": "+1+1", "값": None}])
        book = load_workbook(BytesIO(content))
        self.assertEqual(book.active["B2"].data_type, "s")
        self.assertEqual(book.active["A2"].value, "000000001")
        self.assertEqual(book.active["C2"].value, 0)
        self.assertIsNone(book.active["C3"].value)

    def test_summary_refresh_preserves_details(self):
        rows = [{"mkey": "1", "napgi_compact": date, "gubun": "정기분", "bugwa_amount_won": 900,
                 "total_usage_ton": 12, "sunap_status": "완납"} for date in ("20260430", "20260531")]
        with db.connect() as conn:
            self.assertEqual(import_data.upsert_bill_summaries(conn, rows)["imported"], 2)
            bills = [json.loads(r[0]) for r in conn.execute("SELECT payload FROM bills ORDER BY period")]
        self.assertEqual(bills[0]["납부금액"], 1234)
        self.assertEqual(bills[0]["부과금액"], 900)
        self.assertFalse(bills[0]["summary_only"])
        self.assertTrue(bills[1]["summary_only"])
        self.assertNotIn("상수도_기본료", bills[1])
        self.assertEqual(bills[1]["납기일"], "2026-05-31")
        source = self.root / 'later-details.csv'
        self.write_csv(source, [{'mkey': '000000001', 'napgi': '2026-05', 'gubun': '정기분',
                                'pay_amount_won': '', 'total_usage': 0}])
        with db.connect() as conn:
            import_data.import_details(conn, source)
            bill = json.loads(conn.execute("SELECT payload FROM bills WHERE period='2026-05'").fetchone()[0])
        self.assertFalse(bill['summary_only'])
        self.assertFalse(bill['detail_available'])
        self.assertIsNone(bill['사용량'])
        self.assertEqual(bill['총사용량'], 0)
        self.assertIsNone(catalog.bill_usage(bill))
        with db.connect() as conn:
            import_data.upsert_bill_summaries(conn, [rows[1]])
            record = conn.execute("SELECT payload,source FROM bills WHERE period='2026-05'").fetchone()
            bill = json.loads(record['payload'])
        self.assertEqual(record['source'], 'details_csv')
        self.assertFalse(bill['summary_only'])
        self.assertIsNone(bill['사용량'])
        self.assertEqual(bill['총사용량'], 12)
        self.assertIsNone(catalog.bill_usage(bill))

    def test_live_details_update_values_without_inventing_water_from_groundwater(self):
        rows = [{'mkey': '000000001', 'napgi': '2026-04-30', 'gubun': '정기분', 'bugwa_amount_won': 999,
                 'detail_source': 'i121_public_detail', 'details': {'납부금액': 0, '차감금액': None,
                 '지하수사용량': 4384, '사용량': None, '총사용량': 0, 'periodEnd': '2026-04-20'}}]
        with db.connect() as conn:
            import_data.upsert_bill_summaries(conn, rows)
            record = conn.execute('SELECT payload,source FROM bills').fetchone()
            bill = json.loads(record['payload'])
            self.assertEqual(record['source'], 'i121_public_detail')
            self.assertEqual(bill['납부금액'], 0)
            self.assertEqual(bill['사용량'], 10)  # Missing remote detail does not erase historical detail.
            self.assertEqual(bill['지하수사용량'], 4384)
            self.assertEqual(bill['총사용량'], 0)
            self.assertEqual(bill['periodStart'], '2026-02-01')
            self.assertEqual(bill['periodEnd'], '2026-04-20')
            self.assertTrue(bill['detail_available'])
            self.assertFalse(bill['summary_only'])
            self.assertNotIn('차감금액', bill)

    def test_csv_detail_json_matches_live_dict_and_rejects_invalid_payloads(self):
        details = {'납부금액': 0, '지하수사용량': 4384, '총사용량': 0, '정기검침일': 15, '납부방법': '자동이체'}
        row = {'mkey': '000000001', 'napgi': '2026-04-30', 'gubun': '정기분',
               'detail_source': 'i121_public_detail'}
        with db.connect() as conn:
            for value in (details, json.dumps(details, ensure_ascii=False)):
                import_data.upsert_bill_summaries(conn, [row | {'details': value}])
                bill = json.loads(conn.execute('SELECT payload FROM bills').fetchone()[0])
                self.assertEqual({key: bill[key] for key in details}, details)
            original = dict(conn.execute('SELECT * FROM bills').fetchone())
            for value in ('{bad', '[]', 'null', '0', [], False):
                with self.subTest(value=value), self.assertRaisesRegex(ValueError, '청구 상세 자료'):
                    import_data.upsert_bill_summaries(conn, [row | {'details': value, 'bugwa_amount_won': 9999}])
                self.assertEqual(dict(conn.execute('SELECT * FROM bills').fetchone()), original)
            import_data.upsert_bill_summaries(conn, [row | {'details': ''}])
            bill = json.loads(conn.execute('SELECT payload FROM bills').fetchone()[0])
            self.assertEqual({key: bill[key] for key in details}, details)

    def test_notice_migration_preserves_legacy_ids_and_distinct_supplementary_bills(self):
        with db.connect() as conn:
            original = dict(conn.execute('SELECT * FROM bills').fetchone())
            conn.executescript('''CREATE TABLE old_bills(
                id TEXT PRIMARY KEY,meter_id TEXT NOT NULL REFERENCES meters(id),period TEXT NOT NULL,
                gubun TEXT NOT NULL DEFAULT '정기분',payload TEXT NOT NULL,source TEXT NOT NULL,
                UNIQUE(meter_id,period,gubun));
                INSERT INTO old_bills SELECT id,meter_id,period,gubun,payload,source FROM bills;
                DROP TABLE bills;
                ALTER TABLE old_bills RENAME TO bills;''')
        db.init_db()
        db.init_db()
        rows = [{'mkey': '000000001', 'napgi': '2026-04-08', 'gubun': '수시분',
                 'notice_number': notice, 'bugwa_amount_won': 0} for notice in ('900000001', '900000002')]
        with db.connect() as conn:
            self.assertEqual(dict(conn.execute('SELECT * FROM bills').fetchone()), original)
            self.assertEqual(list(conn.execute('PRAGMA foreign_key_check')), [])
            self.assertEqual(import_data.upsert_bill_summaries(conn, rows)['imported'], 2)
            import_data.upsert_bill_summaries(conn, rows)
            self.assertEqual(conn.execute('SELECT count(*) FROM bills').fetchone()[0], 3)
            self.assertEqual({r[0] for r in conn.execute('SELECT notice_number FROM bills')}, {'', '900000001', '900000002'})
            self.assertEqual(conn.execute('SELECT payload FROM bills WHERE id=?', (original['id'],)).fetchone()[0], original['payload'])
        notices = catalog.payload(['가사업소'])['bills']['000000001']['bills']
        self.assertEqual(len({r['id'] for r in notices}), 3)
        self.assertEqual({r['notice_number'] for r in notices}, {'', '900000001', '900000002'})

    def test_collection_seed_runs_once_and_billing_mode_hides_retained_history(self):
        self.write_csv(self.seed / 'meter_collection.csv', [
            {'customer_number': '000000001', 'arisu_customer_name': '지하철 가역', 'daily_enabled': '0'},
            {'customer_number': '000000002', 'arisu_customer_name': '지하철 나역', 'daily_enabled': '1'},
            {'customer_number': '000000003', 'arisu_customer_name': '등록되지 않은 계약', 'daily_enabled': '1'},
        ])
        with db.connect() as conn:
            catalog.bootstrap(conn)
            meters = {m['id']: m for m in catalog.list_meters(conn)}
            self.assertEqual(len(meters), 2)
            self.assertEqual(meters['000000001']['daily_enabled'], 0)
            self.assertEqual(meters['000000001']['metadata']['arisu_customer_name'], '지하철 가역')
            self.assertEqual(meters['000000002']['daily_enabled'], 1)
        daily = {'000000001': {'usage': [{'date': '2026-09-01', 'value': 0}], 'ridership': []}}
        risk = {'000000001': [{'date': '2026-09-01', 'severity': '경고'}]}
        with patch.object(catalog, 'data_sources', return_value=(daily, risk)):
            payload = catalog.payload()
            self.assertEqual(payload['status']['counts'], {'daily': 0, 'billing_only': 1, 'pending': 1})
            self.assertEqual(payload['daily'], {})
            self.assertEqual(payload['risk'], {})
            self.assertTrue(daily['000000001']['usage'])
        with db.connect() as conn:
            conn.execute("UPDATE meters SET daily_enabled=1,metadata=? WHERE id='000000001'", (json.dumps({'arisu_customer_name': '관리자가 수정한 명칭'}),))
            catalog.bootstrap(conn)
            meter = next(m for m in catalog.list_meters(conn) if m['id'] == '000000001')
            self.assertEqual(meter['daily_enabled'], 1)
            self.assertEqual(meter['metadata']['arisu_customer_name'], '관리자가 수정한 명칭')

    def test_review_uses_presence_column_and_full_unmasked_customer_name(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '수도요금 고지서 정보'
        sheet.append(['고객 번호', '역명', '고지서상 성명', '용도 표기', '유무'])
        sheet.append(['000000001', '가역', '지하철\n 가역', '시민용', None])
        sheet.append(['000000002', '나역', '서울교통공사', '직원용', 'o'])
        path = self.root / 'review.xlsx'
        workbook.save(path)
        with db.connect() as conn:
            self.assertEqual(import_data.import_review(conn, path)['imported'], 2)
            meters = {m['id']: m for m in catalog.list_meters(conn)}
            self.assertEqual(meters['000000001']['daily_enabled'], 0)
            self.assertEqual(meters['000000001']['purpose'], '시민용')
            self.assertEqual(meters['000000001']['metadata']['arisu_customer_name'], '지하철 가역')
            self.assertEqual(meters['000000002']['daily_enabled'], 1)
            self.assertEqual(meters['000000002']['purpose'], '직원용')


if __name__ == "__main__":
    unittest.main()
