"""End-to-end authorization, registry and download checks without external credentials."""
import io
import json
import tempfile
import threading
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import Mock, patch
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from back.api import collection, config, data_access, db, main, mailer


class ApplicationTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.patches = [patch.object(config, 'APP_DB_PATH', root/'test.sqlite3'), patch.object(config, 'SEED_DIR', root),
                        patch.object(config, 'ADMIN_EMAIL', 'chief@example.com'), patch.object(config, 'ADMIN_PASSWORD', 'LocalTestPass!234'),
                        patch.object(config, 'COOKIE_SECURE', False), patch.object(config, 'OPENAI_API_KEY', ''), patch.object(config, 'TODAY_OVERRIDE', ''),
                        patch.object(config, 'COLLECTION_WORKER_ENABLED', False)]
        for p in self.patches:
            p.start()
        self.client = TestClient(main.app, base_url='http://localhost')
        self.client.__enter__()
        with db.connect() as c:
            for office, cid in [('동부', '000000001'), ('서부', '000000002')]:
                c.execute('INSERT INTO offices VALUES(?,?)', (office, office))
                c.execute('INSERT INTO stations VALUES(?,?,?,?)', (office+'역', office+'역', 10, 20))
                c.execute('INSERT INTO meters(id,provider,customer_number,station_id,office_id,display_name,line,daily_enabled,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)', (cid, 'arisu', cid, office+'역', office, office+'역', '2', 1, db.now(), db.now()))
                bill = {'ym': '2026-08', '사용량': 32, '납부금액': 123000}
                c.execute('INSERT INTO bills(id,meter_id,period,gubun,payload,source) VALUES(?,?,?,?,?,?)', (cid+':2026-08:정기분', cid, '2026-08', '정기분', json.dumps(bill), 'test'))
        risk = {cid: [{'date': '2026-08-31', 'actual': 20, 'pred': 10, 'residual': 10, 'z': 4, 'severity': '경고', 'dir': '증가', 'err': False}] for cid in ('000000001','000000002')}
        daily = {cid: {'usage': [{'date': '2026-08-31', 'value': 20}], 'ridership': []} for cid in risk}
        self.source_patch = patch('back.api.catalog.data_sources', return_value=(daily, risk))
        self.source_patch.start()

    def tearDown(self):
        self.source_patch.stop()
        self.client.__exit__(None, None, None)
        for p in reversed(self.patches):
            p.stop()
        self.temp.cleanup()

    def login(self, client=None, email='chief@example.com', password='LocalTestPass!234'):
        client = client or self.client
        response = client.post('/api/auth/login', json={'email': email, 'password': password})
        self.assertEqual(response.status_code, 200, response.text)
        client.headers['X-CSRF-Token'] = response.json()['csrf_token']
        if response.json()['user']['must_change_password']:
            self.assertEqual(client.get('/api/data').status_code, 403)
            changed = client.post('/api/auth/password', json={'current_password': password, 'new_password': password + 'New'})
            self.assertEqual(changed.status_code, 200, changed.text)
        return response.json()['user']

    def create_manager(self, email='staff@example.com', office='동부'):
        response = self.client.post('/api/users', json={'email': email, 'password': 'ManagerPass!234', 'name': '담당자', 'role': 'office_admin', 'office_ids': [office], 'active': True})
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    def test_scoped_session_registry_documents_and_alert(self):
        for path in ('/api/data','/api/meters','/api/users','/api/summary','/api/export.xlsx','/api/bills/000000001:2026-08:정기분/pdf'):
            self.assertEqual(self.client.get(path).status_code, 401, path)
        root = self.login()
        self.assertIn('HttpOnly', self.client.cookies.jar._cookies['localhost.local']['/']['water_session']._rest)
        self.create_manager()
        self.create_manager('staff2@example.com')
        blocked = self.client.post('/api/users', json={'email': 'third@example.com', 'password': 'ManagerPass!234','name': '세번째','role':'office_admin','office_ids':['동부']})
        self.assertEqual(blocked.status_code, 409)
        self.assertEqual(self.client.patch(f"/api/users/{root['id']}", json={'active': False}).status_code, 409)
        manager = TestClient(main.app, base_url='http://localhost')
        staff = self.login(manager, 'staff@example.com', 'ManagerPass!234')
        self.assertEqual(manager.get('/api/users').status_code, 403)
        payload = manager.get('/api/data').json()
        self.assertEqual(set(payload['stations']), {'000000001'})
        self.assertEqual(set(payload['risk']), {'000000001'})
        self.assertEqual(len(manager.get('/api/summary').json()['items']), 1)
        for path in ('/api/bills/000000002:2026-08:정기분/pdf', '/api/export.xlsx?meter_id=000000002'):
            self.assertEqual(manager.get(path).status_code, 404)
        self.assertEqual(manager.get('/api/export.xlsx?office_id=서부').status_code, 403)
        self.assertEqual(manager.post('/api/analyze',json={'meter_id':'000000002'}).status_code, 404)
        self.assertEqual(manager.post('/api/analyze',json={'meter_id':'000000001','date':'2025-01-01'}).status_code, 404)
        self.assertEqual(manager.get('/api/summary?date=bad').status_code, 422)
        self.assertEqual(manager.patch('/api/meters/000000001',json={'office_id':'서부'}).status_code, 403)
        saved = manager.post('/api/meters',json={'provider':'arisu','customer_number':'000000003','station_name':'신규역','office_id':'동부','display_name':'신규역 직원용','line':'2'})
        self.assertEqual(saved.status_code, 201, saved.text)
        self.assertEqual(manager.get('/api/data').json()['bills']['000000003']['bills'], [])
        self.assertEqual(manager.post('/api/meters',json={'customer_number':'000000003','station_name':'신규역','office_id':'동부','display_name':'중복'}).status_code, 409)
        self.assertEqual(manager.patch('/api/meters/000000003',json={'customer_number':'000000004'}).status_code, 422)
        pdf = manager.get('/api/bills/000000001:2026-08:정기분/pdf?download=true')
        self.assertEqual(pdf.status_code, 200)
        self.assertIn('attachment',pdf.headers['content-disposition'])
        self.assertIn('123,000', PdfReader(io.BytesIO(pdf.content)).pages[0].extract_text())
        xlsx = manager.get('/api/export.xlsx?kind=bills')
        rows = list(load_workbook(io.BytesIO(xlsx.content)).active.values)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], '000000001')
        stats = manager.get('/api/stats').json()['rows']
        self.assertEqual(stats[0]['billed_won'],123000)
        self.assertEqual(stats[0]['usage_ton'],20)
        with patch.object(mailer,'send_mail',return_value={'sent':True,'to':['staff@example.com']}) as send:
            response = manager.post('/api/alert',json={'meter_id':'000000001','date':'2026-08-31'})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(set(send.call_args.args[2].split(',')),{'staff@example.com','staff2@example.com'})
            self.assertEqual(manager.post('/api/alert',json={'meter_id':'000000001','date':'2026-08-31'}).status_code,429)
            self.assertEqual(manager.post('/api/alert',json={'meter_id':'000000001','to':'outside@example.com'}).status_code,422)
        self.client.patch(f"/api/users/{staff['id']}",json={'active':False})
        self.assertEqual(manager.get('/api/data').status_code,401)
        manager.close()

    def test_csrf_password_session_and_login_limit(self):
        self.login()
        csrf = self.client.headers.pop('X-CSRF-Token')
        self.assertEqual(self.client.post('/api/auth/logout').status_code,403)
        self.client.headers['X-CSRF-Token'] = csrf
        self.assertEqual(self.client.post('/api/auth/logout',headers={'Origin':'https://outsider.invalid'}).status_code,403)
        old_session = TestClient(main.app, base_url='http://localhost')
        self.login(old_session)
        changed = self.client.post('/api/auth/password',json={'current_password':'LocalTestPass!234','new_password':'ChangedPass!234'})
        self.assertEqual(changed.status_code,200)
        self.assertEqual(old_session.get('/api/data').status_code,401)
        self.assertEqual(self.client.post('/api/auth/logout').status_code,200)
        self.assertEqual(self.client.get('/api/data').status_code,401)
        for _ in range(5):
            self.assertEqual(self.client.post('/api/auth/login',json={'email':'chief@example.com','password':'incorrectpass'}).status_code,401)
        self.assertEqual(self.client.post('/api/auth/login',json={'email':'chief@example.com','password':'ChangedPass!234'}).status_code,429)
        old_session.close()

    def test_other_provider_cannot_alias_foreign_arisu_customer(self):
        self.login()
        self.create_manager()
        manager = TestClient(main.app, base_url='http://localhost')
        self.login(manager, 'staff@example.com', 'ManagerPass!234')
        created = manager.post('/api/meters', json={
            'provider': 'other', 'customer_number': '000000002', 'station_name': '기타공급역',
            'office_id': '동부', 'display_name': '기타 공급 계약', 'daily_enabled': False,
        })
        self.assertEqual(created.status_code, 201, created.text)
        alias_id = created.json()['id']
        _, risk = main.catalog.data_sources()
        risk['000000002'].append(risk['000000002'][0] | {'date': '2026-09-01'})
        scoped = manager.get('/api/data').json()
        self.assertNotIn(alias_id, scoped['risk'])
        self.assertEqual(manager.get('/api/summary').json()['기준일'], '2026-08-31')
        for route in ('summary', 'anomalies'):
            response = manager.get(f'/api/{route}?date=2026-08-31').json()
            self.assertEqual([item['meter_id'] for item in response['items']], ['000000001'])
        for route in ('analyze', 'alert'):
            response = manager.post(f'/api/{route}', json={'meter_id': alias_id, 'date': '2026-09-01'})
            self.assertEqual(response.status_code, 404, response.text)
        self.assertEqual(data_access.station_history(alias_id, '2026-09-01', ['동부']), [])
        manager.close()

    def test_collection_permissions_queue_and_recoverable_deletion(self):
        self.assertEqual(self.client.get('/api/collection').status_code, 401)
        self.login()
        self.create_manager()
        manager = TestClient(main.app, base_url='http://localhost')
        self.login(manager, 'staff@example.com', 'ManagerPass!234')
        self.assertEqual(manager.post('/api/collection', json={'meter_id': '000000002'}).status_code, 404)
        self.assertEqual(manager.delete('/api/meters/000000002').status_code, 404)
        self.assertEqual(manager.post('/api/meters/000000002/restore').status_code, 404)
        self.assertEqual(manager.patch('/api/collection/settings', json={'enabled': False, 'hour': 9}).status_code, 403)
        self.assertEqual(self.client.patch('/api/collection/settings', json={'enabled': True, 'hour': 24}).status_code, 422)
        queued = self.client.post('/api/collection', json={})
        self.assertEqual(queued.status_code, 202, queued.text)
        self.assertEqual(queued.json()['total'], 2)
        self.assertEqual(manager.post('/api/collection', json={}).status_code, 409)
        state = manager.get('/api/collection').json()
        self.assertEqual(state['running']['total'], 1)
        self.assertEqual([m['meter_id'] for m in state['meters']], ['000000001'])
        self.assertNotIn('000000002', json.dumps(state))
        # Deletion withdraws queued work and hides all history until restored.
        self.assertEqual(manager.delete('/api/meters/000000001').status_code, 200)
        self.assertEqual(manager.get('/api/meters').json(), [])
        archived = manager.get('/api/meters?include_deleted=true').json()
        self.assertEqual(len(archived), 1)
        self.assertTrue(archived[0]['deleted_at'])
        self.assertEqual(manager.get('/api/bills/000000001:2026-08:정기분/pdf').status_code, 404)
        self.assertEqual(manager.post('/api/collection', json={'meter_id': '000000001'}).status_code, 404)
        with db.connect() as c:
            self.assertEqual(c.execute('SELECT status FROM collection_items WHERE meter_id=?', ('000000001',)).fetchone()[0], 'skipped')
        restored = manager.post('/api/meters/000000001/restore')
        self.assertEqual(restored.status_code, 200, restored.text)
        self.assertIsNone(restored.json()['deleted_at'])
        self.assertEqual(manager.get('/api/bills/000000001:2026-08:정기분/pdf').status_code, 200)
        self.assertEqual(manager.get('/api/collection').json()['running']['trigger'], 'restoration')
        # A running request holds a stable meter registration until it completes.
        with db.connect() as c:
            c.execute("UPDATE collection_items SET status='running' WHERE meter_id=? AND status='queued'", ('000000001',))
        self.assertEqual(manager.delete('/api/meters/000000001').status_code, 409)
        self.assertEqual(manager.patch('/api/meters/000000001', json={'active': False}).status_code, 409)
        manager.close()

    def test_daily_schedule_is_once_per_kst_day_and_skips_inactive(self):
        with patch.object(collection, 'configured', return_value={'water': True, 'bills': True}):
            with db.connect() as c:
                c.execute('UPDATE meters SET active=0 WHERE id=?', ('000000002',))
                c.execute('BEGIN IMMEDIATE') if not c.in_transaction else None
                current = datetime(2026, 9, 23, 8, 5, tzinfo=collection.KST)
                collection.schedule_due(c, current)
                self.assertEqual(c.execute('SELECT count(*) FROM collection_jobs').fetchone()[0], 1)
                self.assertEqual(c.execute('SELECT count(*) FROM collection_items').fetchone()[0], 1)
                c.execute("UPDATE collection_jobs SET status='success'")
                c.execute("UPDATE collection_items SET status='success'")
                collection.schedule_due(c, current)
                self.assertEqual(c.execute('SELECT count(*) FROM collection_jobs').fetchone()[0], 1)
                collection.schedule_due(c, datetime(2026, 9, 24, 7, tzinfo=collection.KST))
                self.assertEqual(c.execute('SELECT count(*) FROM collection_jobs').fetchone()[0], 1)
                collection.schedule_due(c, datetime(2026, 9, 24, 8, tzinfo=collection.KST))
                self.assertEqual(c.execute('SELECT count(*) FROM collection_jobs').fetchone()[0], 2)

    def test_collection_worker_keeps_partial_results_and_no_false_connection(self):
        self.login()
        job = self.client.post('/api/collection', json={}).json()
        with db.connect() as c:
            c.execute("UPDATE collection_jobs SET status='running' WHERE id=?", (job['id'],))
        results = [dict(status='success', message='새 청구서 1건', daily_rows=0, bill_rows=1, connection_verified=True),
                   dict(status='empty', message='조회 기간에 자료 없음', daily_rows=0, bill_rows=0, connection_verified=False)]
        session = Mock()
        with patch('back.pipelines.incremental.collect_meter', side_effect=results) as collect, patch('back.scripts.billing_etl.i121_crawler.auth.session_from_env', return_value=session) as login:
            collection.run_job(job['id'], threading.Event())
            self.assertEqual(collect.call_count, 2)
            login.assert_called_once()
            session.close.assert_called_once()
        state = self.client.get('/api/collection').json()
        self.assertIsNone(state['running'])
        self.assertEqual(state['latest']['completed'], 2)
        self.assertEqual([m['connection_verified'] for m in state['meters']], [True, False])
        self.assertIsNone(state['meters'][1]['last_success_at'])
        with db.connect() as c:
            self.assertEqual(c.execute('SELECT count(*) FROM bills').fetchone()[0], 2)

    def test_collection_session_failure_is_reported_once_and_preserves_last_success(self):
        from requests import Response
        from requests.exceptions import HTTPError
        self.login()
        previous = self.client.post('/api/collection', json={}).json()
        collection._run_job(previous['id'], threading.Event(), lambda *args, **kwargs:
                            {'status': 'success', 'message': '자료 확인', 'connection_verified': True})
        before = self.client.get('/api/collection').json()['meters']
        job = self.client.post('/api/collection', json={}).json()
        response = Response()
        response.status_code = 404
        error = HTTPError('secret must never reach clients', response=response)
        with patch('back.scripts.billing_etl.i121_crawler.auth.collection_session', side_effect=error) as login, \
             patch('back.pipelines.incremental.collect_meter') as collect, \
             patch.object(collection.logger, 'error') as log:
            collection.run_job(job['id'], threading.Event())
            login.assert_called_once()
            collect.assert_not_called()
            log.assert_called_once()
        state = self.client.get('/api/collection').json()
        self.assertIsNone(state['running'])
        self.assertEqual(state['latest']['status'], 'failed')
        self.assertEqual(state['latest']['completed'], 2)
        self.assertIn('실패 2개', state['latest']['message'])
        self.assertIn('HTTP 404', state['latest']['message'])
        self.assertNotIn('확인 완료', state['latest']['message'])
        self.assertNotIn('secret', json.dumps(state))
        self.assertTrue(all('공통 연결 단계' in item['message'] for item in state['latest']['items']))
        self.assertEqual([m['last_success_at'] for m in state['meters']], [m['last_success_at'] for m in before])
        with db.connect() as c:
            self.assertEqual(c.execute('SELECT count(*) FROM bills').fetchone()[0], 2)

    def test_collection_failure_messages_distinguish_causes(self):
        from requests import Response
        from requests.exceptions import ConnectionError, HTTPError, Timeout
        from back.scripts.billing_etl.i121_crawler.auth import ConfigurationError, LoginError
        cases = [(ConfigurationError('private'), '설정'), (LoginError('private'), '로그인'),
                 (Timeout('private'), '응답 시간'), (ConnectionError('private'), '네트워크')]
        for code in (401, 403, 404, 429, 503):
            response = Response()
            response.status_code = code
            cases.append((HTTPError('private', response=response), f'HTTP {code}'))
        for error, expected in cases:
            with self.subTest(error=type(error).__name__, expected=expected):
                result = collection.failure_result(error)
                self.assertIn(expected, result['message'])
                self.assertNotIn('private', result['message'])
                self.assertIn('기존 자료는 보존', result['message'])

    def test_session_failure_leaves_restored_meter_in_its_new_queue(self):
        from back.scripts.billing_etl.i121_crawler.auth import LoginError
        self.login()
        job = self.client.post('/api/collection', json={}).json()

        def failed_login(*args):
            self.assertEqual(self.client.delete('/api/meters/000000002').status_code, 200)
            self.assertEqual(self.client.post('/api/meters/000000002/restore').status_code, 200)
            raise LoginError('session setup failed')

        with patch('back.scripts.billing_etl.i121_crawler.auth.collection_session', side_effect=failed_login):
            collection.run_job(job['id'], threading.Event())
        with db.connect() as c:
            states = [r[0] for r in c.execute('SELECT status FROM collection_items WHERE meter_id=? ORDER BY job_id', ('000000002',))]
            self.assertEqual(states, ['skipped', 'queued'])

    def test_deleting_and_restoring_queued_meter_cannot_run_it_twice(self):
        self.login()
        job = self.client.post('/api/collection', json={}).json()
        called = []

        def collect(meter, **kwargs):
            called.append(meter['id'])
            self.assertEqual(self.client.delete('/api/meters/000000002').status_code, 200)
            self.assertEqual(self.client.post('/api/meters/000000002/restore').status_code, 200)
            return {'status': 'empty', 'message': '조회 자료 없음'}

        collection._run_job(job['id'], threading.Event(), collect)
        self.assertEqual(called, ['000000001'])
        with db.connect() as c:
            states = [r[0] for r in c.execute('SELECT status FROM collection_items WHERE meter_id=? ORDER BY job_id', ('000000002',))]
            self.assertEqual(states, ['skipped', 'queued'])

    def test_risk_requires_observed_daily_water(self):
        self.login()
        daily, risk = main.catalog.data_sources()
        daily['000000001']['usage'][0]['value'] = 0
        daily['000000002'] = {'usage': [], 'ridership': [{'date': '2026-09-01', 'value': 100}]}
        risk['000000002'].append(risk['000000002'][0] | {'date': '2026-09-01'})
        for enabled, expected in [(0, 'billing_only'), (1, 'pending')]:
            with db.connect() as c:
                c.execute("UPDATE meters SET daily_enabled=? WHERE id='000000002'", (enabled,))
            payload = self.client.get('/api/data').json()
            modes = {m['id']: m['data_mode'] for m in payload['meters']}
            self.assertEqual(modes, {'000000001': 'daily', '000000002': expected})
            self.assertEqual(payload['stations']['000000002']['data_mode'], expected)
            self.assertEqual(payload['status']['counts'], {'daily': 1, 'billing_only': 1-enabled, 'pending': enabled})
            self.assertEqual(set(payload['risk']), {'000000001'})
            self.assertEqual(payload['status']['latest_risk'], '2026-08-31')
            self.assertEqual(len(payload['bills']['000000002']['bills']), 1)
            summary = self.client.get('/api/summary').json()
            self.assertEqual(summary['기준일'], '2026-08-31')
            self.assertEqual(summary['counts']['경고'], 1)
            self.assertEqual([item['meter_id'] for item in summary['items']], ['000000001'])
            self.assertEqual(self.client.get('/api/anomalies?date=2026-09-01').json()['items'], [])
            for route in ('analyze', 'alert'):
                self.assertEqual(self.client.post(f'/api/{route}', json={'meter_id': '000000002', 'date': '2026-09-01'}).status_code, 404)
            self.assertEqual(data_access.station_history('000000002', '2026-09-01'), [])
            # A collected row without an actual value still is not daily water data.
            daily['000000002']['usage'] = [{'date': '2026-09-01', 'value': None}]

    def test_invalid_observation_is_separate_from_alerts_and_cannot_trigger_external_actions(self):
        self.login()
        _, risk = main.catalog.data_sources()
        risk['000000001'][0].update(actual=-1, residual=-11, z=-5, err=True)
        risk['000000002'][0].update(actual=3186, residual=3176, z=8, err=False)
        summary = self.client.get('/api/summary').json()
        self.assertEqual(summary['counts'], {'경고': 1, '주의': 0, '자료 확인': 1, '총': 2, '분석': 1})
        self.assertIn('자료 확인 1건', summary['headline'])
        invalid = next(item for item in summary['items'] if item['meter_id'] == '000000001')
        self.assertEqual(invalid['심각도'], '자료 확인')
        self.assertTrue(invalid['likely_data_error'])
        anomaly = next(item for item in self.client.get('/api/anomalies').json()['items'] if item['meter_id'] == '000000001')
        self.assertEqual(anomaly['심각도'], '자료 확인')
        with patch.object(main.agent, 'analyze_cause') as analyze, patch.object(mailer, 'send_mail') as send:
            for route in ('analyze', 'alert'):
                result = self.client.post('/api/' + route, json={'meter_id': '000000001', 'date': '2026-08-31'})
                self.assertEqual(result.status_code, 422, result.text)
                self.assertIn('원자료', result.json()['detail'])
            analyze.assert_not_called()
            send.assert_not_called()
        risk['000000002'][0]['err'] = True
        only_errors = self.client.get('/api/summary').json()
        self.assertEqual(only_errors['counts']['분석'], 0)
        self.assertEqual(only_errors['counts']['자료 확인'], 2)
        self.assertEqual(only_errors['actions'], ['자료 확인 항목은 원자료를 확인한 뒤 다시 분석하세요.'])

    def test_inactive_meter_history_does_not_move_dashboard_reference_dates(self):
        self.login()
        daily, risk = main.catalog.data_sources()
        daily['000000001']['ridership'] = [{'date': '2026-08-30', 'value': 10}]
        daily['000000002']['usage'][0]['date'] = '2026-09-01'
        daily['000000002']['ridership'] = [{'date': '2026-09-01', 'value': 20}]
        risk['000000002'][0]['date'] = '2026-09-01'
        with db.connect() as c:
            c.execute('UPDATE meters SET active=0 WHERE id=?', ('000000002',))
        payload = self.client.get('/api/data').json()
        self.assertEqual(payload['status']['reference_date'], '2026-08-31')
        self.assertEqual(payload['status']['latest_usage'], '2026-08-31')
        self.assertEqual(payload['status']['latest_ridership'], '2026-08-30')
        self.assertEqual(payload['status']['latest_risk'], '2026-08-31')
        self.assertIn('000000002', payload['daily'])
        self.assertIn('000000002', payload['risk'])
        summary = self.client.get('/api/summary', params={'date': payload['status']['reference_date']}).json()
        self.assertEqual(summary['counts']['경고'], 1)

    def test_detailed_groundwater_bill_does_not_become_zero_water_usage(self):
        self.login()
        with db.connect() as c:
            c.execute('UPDATE bills SET payload=? WHERE meter_id=?', (json.dumps({
                '사용량': None, '총사용량': 0, '지하수사용량': 4384,
                '납부금액': 2104320, 'detail_available': True,
            }), '000000001'))
        row = self.client.get('/api/stats?meter_id=000000001').json()['rows'][0]
        self.assertIsNone(row['billed_usage_ton'])
        self.assertEqual(row['billed_won'], 2104320)

    def test_imported_missing_usage_matches_statistics_export_and_pdf(self):
        from back.api.import_data import import_details
        self.login()
        source = Path(self.temp.name) / 'missing-usage.csv'
        source.write_text('mkey,napgi,gubun,pay_amount_won,total_usage\n000000001,2026-09,정기분,,0\n', encoding='utf-8')
        with db.connect() as c:
            self.assertEqual(import_details(c, source)['imported'], 1)
        query = {'meter_id': '000000001', 'start': '2026-09-01', 'end': '2026-09-30'}
        row = self.client.get('/api/stats', params=query).json()['rows'][0]
        self.assertIsNone(row['billed_usage_ton'])
        response = self.client.get('/api/export.xlsx', params=query | {'kind': 'bills'})
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        values = list(workbook.active.values)
        workbook.close()
        exported = dict(zip(values[0], values[1]))
        self.assertIsNone(exported['사용량'])
        self.assertEqual(exported['총사용량'], 0)
        self.assertIsNone(exported['집계사용량_톤'])
        pdf = self.client.get('/api/bills/000000001:2026-09:정기분/pdf')
        text = ''.join(PdfReader(io.BytesIO(pdf.content)).pages[0].extract_text().split())
        self.assertIn('상·하수도사용량(톤)-', text)

    def test_invalid_model_response_uses_safe_failure_instead_of_react_objects(self):
        self.login()
        valid = {'primary_cause': '원자료와 운영 현황 확인이 필요합니다.', 'confidence': '낮음',
                 'recommendation': '사용량 원자료를 확인하세요.', 'reasons': ['관측값의 변동'],
                 'events': [{'title': '관련 정보', 'date': '2026-08-31', 'source': 'https://example.com/info'}]}
        service = Mock()
        with patch.object(config, 'OPENAI_API_KEY', 'mock'), patch.object(main.agent, 'OpenAI', return_value=service):
            for invalid in ({'confidence': {'level': 'low'}}, {'recommendation': ['원자료 확인']},
                            {'reasons': [{'text': '자료 확인'}]}, {'is_calendar_effect': 'false'},
                            {'events': [{'title': {'name': '행사'}, 'source': 'https://example.com'}]}):
                service.responses.create.return_value.output_text = json.dumps(valid | invalid, ensure_ascii=False)
                response = self.client.post('/api/analyze', json={'meter_id': '000000001'}).json()
                self.assertEqual(response['generated_by'], 'none')
                self.assertIn('error', response['analysis'])
            service.responses.create.return_value.output_text = json.dumps(valid, ensure_ascii=False)
            response = self.client.post('/api/analyze', json={'meter_id': '000000001'}).json()
            self.assertEqual(response['analysis'], valid)

    def test_collection_fields_are_validated_and_existing_metadata_is_preserved(self):
        self.login()
        with db.connect() as c:
            c.execute('UPDATE meters SET metadata=? WHERE id=?', (json.dumps({'review_notes': {'note': 'keep'}}), '000000001'))
        result = self.client.patch('/api/meters/000000001', json={'daily_enabled': False,
                'metadata': {'arisu_customer_name': '  지하철 가역  ', '계량기번호': '123'}})
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['metadata']['arisu_customer_name'], '지하철 가역')
        self.assertEqual(result.json()['metadata']['review_notes'], {'note': 'keep'})
        payload = self.client.get('/api/data').json()
        self.assertEqual(payload['stations']['000000001']['data_mode'], 'billing_only')
        self.assertNotIn('000000001', payload['daily'])
        self.assertNotIn('000000001', payload['risk'])
        self.assertEqual(len(payload['bills']['000000001']['bills']), 1)
        workbook = load_workbook(io.BytesIO(self.client.get('/api/export.xlsx?kind=usage').content), read_only=True)
        exported = list(workbook.active.values)
        workbook.close()
        self.assertEqual(len(exported), 2)
        self.assertEqual(exported[1][exported[0].index('고객번호')], '000000002')
        for invalid in ('가*역', '가＊역', 'bad\nname', 123):
            response = self.client.patch('/api/meters/000000001', json={'metadata': {'arisu_customer_name': invalid}})
            self.assertEqual(response.status_code, 422)


if __name__ == '__main__':
    unittest.main()
